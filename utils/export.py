"""Summary export formatters."""

from __future__ import annotations
from io import BytesIO
from typing import Any
from xml.sax.saxutils import escape


def to_markdown(result: dict[str, Any]) -> str:
    """Create a Markdown document from a summary result."""
    parts = []

    if result.get("title"):
        parts.append(f"# {result['title']}")

    parts.append(str(result.get("summary", "")))

    if result.get("key_points"):
        parts.extend(
            ["", "## Key points", *[f"- {item}" for item in result["key_points"]]]
        )

    if result.get("keywords"):
        parts.extend(["", "**Keywords:** " + ", ".join(result["keywords"])])

    return "\n\n".join(parts).strip() + "\n"


def create_docx(result: dict[str, Any]) -> bytes:
    """Create an in-memory DOCX summary."""
    try:
        from docx import Document
    except ImportError as exc:
        raise RuntimeError("Install python-docx to export DOCX file.") from exc

    document = Document()

    if result.get("title"):
        document.add_heading(str(result["title"]), level=1)
    document.add_paragraph(str(result.get("summary", "")))

    if result.get("key_points"):
        document.add_heading("Key points", level=2)
        for point in result["key_points"]:
            document.add_paragraph(str(point), style="List Bullet")

    if result.get("keywords"):
        document.add_paragraph("Keywords:" + ", ".join(result["keywords"]))

    output = BytesIO()
    document.save(output)

    return output.getvalue()


def create_pdf(result: dict[str, Any]) -> bytes:
    """Create a simple in-memory PDF summary."""
    try:
        from reportlab.lib.pagesizes import letter
        from reportlab.lib.styles import getSampleStyleSheet
        from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer
    except ImportError as exc:
        raise RuntimeError("Install reportlab to export PDF files.") from exc

    output = BytesIO()
    document = SimpleDocTemplate(output, pagesize=letter)
    styles = getSampleStyleSheet()
    content = []

    if result.get("title"):
        content.append(Paragraph(escape(str(result["title"])), styles["Title"]))
        content.append(Spacer(1, 22))

    content.append(
        Paragraph(
            escape(str(result.get("summary", ""))).replace("\n", "<br/>"),
            styles["BodyText"],
        )
    )

    if result.get("key_points"):
        content.extend([Spacer(1, 12), Paragraph("Key points", styles["Heading2"])])
        for point in result["key_points"]:
            content.append(
                Paragraph("&#8226; " + escape(str(point)), styles["BodyText"])
            )

    document.build(content)

    return output.getvalue()
