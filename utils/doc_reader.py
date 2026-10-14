"""DOCX text extraction."""

from __future__ import annotations
from io import BytesIO
from typing import Any

from docx import Document
from utils.helpers import save_uploaded_text


def read_docx(upload: Any) -> str:
    """Extract paragraph and table text from a DOCX upload."""
    raw = save_uploaded_text(upload)

    try:
        document = Document(BytesIO(raw))
        sections = [
            paragraph.text.strip()
            for paragraph in document.paragraphs
            if paragraph.text.strip()
        ]

        for table in document.tables:
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if cells:
                    sections.append(" | ".join(cells))
    except Exception as exc:
        raise ValueError(f"Could not read this DOCX file: {exc}") from exc

    text = "\n".join(sections).strip()
    if not text:
        raise ValueError("The DOCX file does not contain readable text.")

    return text
