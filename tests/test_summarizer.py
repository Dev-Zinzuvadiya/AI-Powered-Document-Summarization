"""Unit tests for offline summarization utilities."""

import unittest
from io import BytesIO

from docx import Document

from utils.doc_reader import read_docx
from utils.export import create_docx, create_pdf, to_markdown
from utils.summarizer import summarize
from utils.tokenizer import chunk_text, count_tokens
from utils.web_scraper import fetch_webpage
from utils.youtube_loader import extract_video_id


class UploadedFile(BytesIO):
    def __init__(self, data: bytes, name: str) -> None:
        super().__init__(data)
        self.name = name
        self.size = len(data)

    def getvalue(self) -> bytes:
        self.seek(0)
        return super().getvalue()


class SummarizerTests(unittest.TestCase):
    def test_extractive_summary_reports_metrics(self) -> None:
        source = (
            "Solar power is a renewable energy source. Solar panels convert sunlight to electricity. "
            "The panels can reduce household electricity costs. Battery storage can keep excess power. "
            "Many communities are investing in renewable energy."
        )

        result = summarize(source, ratio=0.4, engine="extractive")

        self.assertTrue(result["summary"])
        self.assertEqual(result["engine"], "extractive")
        self.assertGreater(result["source_words"], result["summary_words"])
        self.assertTrue(result["keywords"])

    def test_empty_input_is_rejected(self) -> None:
        with self.assertRaises(ValueError):
            summarize(" ", engine="extractive")

    def test_chunking_respects_limit(self) -> None:
        chunks = list(
            chunk_text(("A sentence with useful detail. " * 100), max_chars=120)
        )
        self.assertGreater(len(chunks), 1)
        self.assertTrue(all(len(chunk) <= 120 for chunk in chunks))

    def test_token_count_returns_positive_estimate(self) -> None:
        self.assertGreater(count_tokens("One short sentence.", 0))

    def test_docx_extracts_paragraphs_and_tables(self) -> None:
        document = Document()
        document.add_paragraph("Quarterly results")

        table = document.add_table(rows=1, cols=1)
        table.cell(0, 0).text = "Revenue increased"

        output = BytesIO()
        document.save(output)

        source = read_docx(UploadedFile(output.getvalue(), "report.docx"))

        self.assertIn("Quarterly results", source)
        self.assertIn("Revenue increased", source)

    def test_export_formats_are_generated(self) -> None:
        result = {
            "title": "Test",
            "summary": "A summary.",
            "key_points": ["Point."],
            "keywords": ["summary"],
        }

        self.assertIn("# Test", to_markdown(result))
        self.assertTrue(create_docx(result).startswith(b"PK"))
        self.assertTrue(create_pdf(result).startswith(b"%PDF"))

    def test_private_web_urls_are_rejected(self) -> None:
        with self.assertRaisesRegex(ValueError, "private network"):
            fetch_webpage("https://127.0.0.1/")

    def test_youtube_url_video_id(self) -> None:
        self.assertEqual(
            extract_video_id("https://youtu.be/abcdefghijk"), "abcdefghijk"
        )

        with self.assertRaises(ValueError):
            extract_video_id("https://example.com/")


if __name__ == "__main__":
    unittest.main()
